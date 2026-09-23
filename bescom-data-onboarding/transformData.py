#!/usr/bin/env python3
"""Convert the BESCOM officer workbook into contacts.json with canonical hierarchy.

Primary record sheets:
  - Field Section
  - Corporate Office
Detailed sheets are indexed to enrich PIN and phone values where possible.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("openpyxl is required: python -m pip install openpyxl") from exc

PLACEHOLDER_PHONES = {"9999999999", "0000000000"}
CORPORATE_MARKERS = ("corporate office", "head office")


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def norm(value: Any) -> str:
    value = unicodedata.normalize("NFKD", text(value)).lower()
    value = value.replace("&", " and ").replace("\u200d", " ")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def compact(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", norm(value))


def clean_date(value: Any) -> str:
    raw = text(value)
    if not raw:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    raw = raw.replace("/", ".").replace("-", ".")
    match = re.search(r"\b(\d{1,2})\.(\d{1,2})\.(\d{2,4})\b", raw)
    if not match:
        return text(value)
    day, month, year = map(int, match.groups())
    if year < 100:
        year += 2000 if year <= 50 else 1900
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return text(value)


def clean_pin(value: Any) -> str:
    raw = text(value)
    return raw[:-2] if raw.endswith(".0") else raw


def extract_phones(value: Any) -> list[str]:
    raw = text(value)
    if not raw:
        return []
    candidates = re.findall(r"(?:\+?91[\s-]?)?(?:0\d{2,4}[\s-]?\d{5,8}|\d[\d\s/-]{7,14}\d)", raw)
    result: list[str] = []
    for candidate in candidates:
        digits = re.sub(r"\D", "", candidate)
        if digits.startswith("91") and len(digits) == 12:
            digits = digits[2:]
        if digits in PLACEHOLDER_PHONES or len(digits) < 9 or len(digits) > 11:
            continue
        if digits not in result:
            result.append(digits)
    return result


def canonical_designation(raw: str) -> str:
    n = norm(raw)
    if "chief controllers" in n or "chief controller" in n:
        return "Chief Controller of Accounts"
    if "deputy controller" in n:
        return "Deputy Controller of Accounts"
    if "accounts officer" in n:
        return "Accounts Officer"
    if "assistant executive engineer" in n:
        return "Assistant Executive Engineer (Civil)" if "civil" in n else "Assistant Executive Engineer"
    if "superintending engineer" in n:
        return "Superintending Engineer"
    if "executive engineer" in n:
        return "Executive Engineer (Civil)" if "civil" in n else "Executive Engineer"
    if "chief engineer" in n:
        return "Chief Engineer"
    if "chief general manager" in n:
        return "Chief General Manager"
    if "assistant general manager" in n:
        return "Assistant General Manager"
    if "deputy general manager" in n:
        return "Deputy General Manager"
    if "general manager" in n:
        return "General Manager"
    if "company secretary" in n:
        return "Company Secretary"
    if "director" in n:
        return "Director"
    return text(raw)


def category_for(designation: str, corporate: bool) -> str:
    if corporate:
        return "Corporate Office"
    n = norm(designation)
    return "Accounts" if any(x in n for x in ("account", "controller", "finance")) else "Engineer"


def build_enrichment_index(workbook) -> dict[str, list[dict[str, str]]]:
    layouts = {
        "SEE": (2, 3, 4, 7), "DCA": (2, 3, 4, 6), "EE(E) ": (2, 3, 4, 6),
        "EE(civil)": (2, 3, 4, None), "AEE(E) ": (4, 5, 6, 10),
        "AO": (3, 4, 5, 8), "AEE(Ele)": (3, 4, 5, 10),
        "Sheet1": (2, 3, 4, 5), "Sheet2": (2, 3, 4, 5),
    }
    index: dict[str, list[dict[str, str]]] = defaultdict(list)
    for sheet_name, (office_col, name_col, pin_col, phone_col) in layouts.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        for row_number in range(1, sheet.max_row + 1):
            name = text(sheet.cell(row_number, name_col).value)
            name_key = norm(name)
            if not name_key or name_key in {"vacant", "name of the officer serving sriyuths", "name of the officer vacant"}:
                continue
            if len(name_key) < 3 or "section" in name_key and not text(sheet.cell(row_number, pin_col).value):
                continue
            phones = extract_phones(sheet.cell(row_number, phone_col).value) if phone_col else []
            entry = {
                "name": name,
                "office": text(sheet.cell(row_number, office_col).value),
                "pin": clean_pin(sheet.cell(row_number, pin_col).value),
                "primaryPhone": phones[0] if phones else "",
                "secondaryPhones": phones[1:],
                "source": f"{sheet_name}!{row_number}",
            }
            if entry["pin"] or entry["primaryPhone"]:
                index[name_key].append(entry)
    return index


def token_score(left: str, right: str) -> int:
    a, b = set(norm(left).split()), set(norm(right).split())
    return len(a & b)


def choose_enrichment(index: dict[str, list[dict[str, str]]], name: str, office: str) -> dict[str, Any]:
    candidates = index.get(norm(name), [])
    if not candidates:
        return {}
    return max(candidates, key=lambda x: (token_score(x["office"], office), bool(x["primaryPhone"]), bool(x["pin"])))


class HierarchyResolver:
    def __init__(self, mapping: dict[str, Any]):
        self.mapping = mapping
        self.circle_to_zone = {
            circle: zone for zone, circles in mapping["zoneCircles"].items() for circle in circles
        }
        self.zone_aliases = self._compile_aliases(mapping["zoneAliases"])
        self.circle_aliases = self._compile_aliases(mapping["circleAliases"])
        self.division_aliases: list[tuple[str, str]] = []
        for division, spec in mapping["divisions"].items():
            for alias in spec["aliases"]:
                self.division_aliases.append((norm(alias), division))
        self.division_aliases.sort(key=lambda pair: len(pair[0]), reverse=True)
        self.code_to_division = {
            code.upper(): division
            for division, codes in mapping.get("sectionCodes", {}).items()
            for code in codes
        }

    @staticmethod
    def _compile_aliases(groups: dict[str, list[str]]) -> list[tuple[str, str]]:
        result = [(norm(alias), canonical) for canonical, aliases in groups.items() for alias in aliases]
        return sorted(result, key=lambda pair: len(pair[0]), reverse=True)

    @staticmethod
    def _contains(haystack: str, needle: str) -> bool:
        if not needle:
            return False
        if re.search(rf"(?:^|\s){re.escape(needle)}(?:$|\s)", haystack):
            return True
        compact_needle = re.sub(r"[^a-z0-9]", "", needle)
        compact_haystack = re.sub(r"[^a-z0-9]", "", haystack)
        return len(compact_needle) >= 6 and compact_needle in compact_haystack

    def _match(self, office_norm: str, aliases: list[tuple[str, str]]) -> str:
        for alias, canonical in aliases:
            if self._contains(office_norm, alias):
                return canonical
        return ""

    def _division_from_code(self, raw_office: str) -> str:
        upper = text(raw_office).upper()
        codes = re.findall(r"(?<![A-Z0-9])([SNEWCK]-?\d{1,2})(?![A-Z0-9])", upper)
        for raw_code in codes:
            code = raw_code.replace("-", "")
            if code in self.code_to_division:
                return self.code_to_division[code]
        return ""

    def resolve(self, office: str, corporate: bool = False) -> dict[str, str]:
        if corporate or any(marker in norm(office) for marker in CORPORATE_MARKERS):
            return {"zone": "", "circle": "", "division": "", "mappingStatus": "not_applicable"}

        office_norm = norm(office)
        division = ""
        for alias, canonical in self.division_aliases:
            if self._contains(office_norm, alias):
                division = canonical
                break
        if not division:
            division = self._division_from_code(office)

        circle = self.mapping["divisions"].get(division, {}).get("circle", "") if division else ""
        explicit_circle = self._match(office_norm, self.circle_aliases)
        if explicit_circle and not circle:
            circle = explicit_circle
        zone = self.circle_to_zone.get(circle, "")

        explicit_zone = self._match(office_norm, self.zone_aliases)
        if explicit_zone and not zone:
            zone = explicit_zone

        if zone and circle and circle not in self.mapping["zoneCircles"].get(zone, []):
            zone = self.circle_to_zone.get(circle, zone)

        if zone and circle:
            status = "resolved"
        elif zone or circle or division:
            status = "partial"
        else:
            status = "unresolved"
        return {"zone": zone, "circle": circle, "division": division, "mappingStatus": status}


def infer_level(office: str, hierarchy: dict[str, str], corporate: bool) -> str:
    n = norm(office)
    if corporate:
        return "Corporate Office"
    if "sub division" in n or "sub div" in n or re.search(r"\b[snewck]\s*\d{1,2}\b", n):
        return "Subdivision"
    if "section" in n and "division" not in n:
        return "Section"
    if hierarchy["division"] or "division" in n or " div " in f" {n} " or "works unit" in n:
        return "Division"
    if hierarchy["circle"] or "circle" in n:
        return "Circle"
    if hierarchy["zone"] or "zone" in n:
        return "Zone"
    return "Office"


def record_from_row(row: tuple[Any, ...], source_sheet: str, source_row: int,
                    resolver: HierarchyResolver, enrichment: dict[str, list[dict[str, str]]],
                    debug: bool, hierarchy_override: dict[str, str] | None = None) -> dict[str, Any] | None:
    corporate = source_sheet == "Corporate Office"
    if corporate:
        serial, name, raw_designation, raw_phone, office, section, charge_date = (list(row) + [None] * 7)[:7]
        office_name = text(section) or text(office) or "Corporate Office"
        address = text(office) or "Corporate Office"
    else:
        serial, name, raw_designation, raw_phone, office, charge_date = (list(row) + [None] * 6)[:6]
        section = ""
        office_name = text(office)
        address = office_name

    name = text(name)
    if not name or norm(name) in {"name", "vacant"}:
        return None

    raw_designation = text(raw_designation)
    designation = canonical_designation(raw_designation)
    hierarchy = hierarchy_override or resolver.resolve(office_name, corporate)
    enriched = choose_enrichment(enrichment, name, office_name)
    workbook_phones = extract_phones(raw_phone)
    phones = workbook_phones if workbook_phones else ([enriched.get("primaryPhone", "")] if enriched.get("primaryPhone") else [])
    for phone in enriched.get("secondaryPhones", []):
        if phone and phone not in phones:
            phones.append(phone)

    serial_text = clean_pin(serial) or str(source_row - 1)
    prefix = "CO" if corporate else "FS"
    record: dict[str, Any] = {
        "id": f"{prefix}-{serial_text}",
        "code": f"{prefix}-{serial_text}",
        "sourceSheet": source_sheet,
        "sourceRow": source_row,
        "level": infer_level(office_name, hierarchy, corporate),
        "zone": hierarchy["zone"],
        "circle": hierarchy["circle"],
        "division": hierarchy["division"],
        "subdivision": "",
        "section": text(section),
        "category": category_for(designation, corporate),
        "status": "Active",
        "officeName": office_name,
        "address": address,
        "dateOfAssumingCharge": clean_date(charge_date),
        "remarks": "",
        "officer": {
            "name": name,
            "designation": designation,
            "rawDesignation": raw_designation,
            "pinNumber": enriched.get("pin", ""),
            "email": "",
            "primaryPhone": phones[0] if phones else "",
            "secondaryPhones": phones[1:] if len(phones) > 1 else [],
            "mobile": phones[0] if phones else ""
        }
    }
    if debug:
        record["mappingStatus"] = hierarchy["mappingStatus"]
        record["enrichmentSource"] = enriched.get("source", "")
    return record


def write_report(path: Path, records: list[dict[str, Any]]) -> None:
    fields = ["id", "sourceSheet", "sourceRow", "name", "designation", "officeName", "zone", "circle", "division", "mappingStatus"]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for record in records:
            status = record.get("mappingStatus") or (
                "not_applicable" if record["category"] == "Corporate Office"
                else "resolved" if record["zone"] and record["circle"]
                else "partial" if record["zone"] or record["circle"] or record["division"]
                else "unresolved"
            )
            if status in {"partial", "unresolved"}:
                writer.writerow({
                    "id": record["id"], "sourceSheet": record["sourceSheet"],
                    "sourceRow": record["sourceRow"], "name": record["officer"]["name"],
                    "designation": record["officer"]["designation"], "officeName": record["officeName"],
                    "zone": record["zone"], "circle": record["circle"], "division": record["division"],
                    "mappingStatus": status,
                })


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert the BESCOM officer workbook to contacts.json")
    parser.add_argument("input", type=Path, help="Input .xlsx workbook")
    parser.add_argument("-o", "--output", type=Path, default=Path("contacts.json"))
    parser.add_argument("-m", "--mappings", type=Path, default=Path(__file__).with_name("hierarchy_mappings.json"))
    parser.add_argument("-r", "--report", type=Path, default=Path("mapping_report.csv"))
    parser.add_argument("--debug", action="store_true", help="Keep mappingStatus/enrichmentSource in contacts.json")
    args = parser.parse_args()

    if not args.input.exists():
        parser.error(f"input workbook not found: {args.input}")
    if not args.mappings.exists():
        parser.error(f"mapping file not found: {args.mappings}")

    mapping = json.loads(args.mappings.read_text(encoding="utf-8"))
    resolver = HierarchyResolver(mapping)
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    required = {"Field Section", "Corporate Office"}
    missing = required - set(workbook.sheetnames)
    if missing:
        raise SystemExit(f"Missing required sheet(s): {', '.join(sorted(missing))}")

    enrichment = build_enrichment_index(workbook)
    records: list[dict[str, Any]] = []
    for sheet_name in ("Field Section", "Corporate Office"):
        sheet = workbook[sheet_name]
        context = {"zone": "", "circle": "", "division": ""}
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            hierarchy_override = None
            if sheet_name == "Field Section":
                serial = clean_pin(row[0] if row else "")
                direct = resolver.resolve(text(row[4] if len(row) > 4 else ""), False)
                try:
                    ordered_detail = int(float(serial)) >= 114
                except (TypeError, ValueError):
                    ordered_detail = False

                if ordered_detail:
                    if direct["division"]:
                        context = {key: direct[key] for key in ("zone", "circle", "division")}
                    elif direct["circle"]:
                        context = {"zone": direct["zone"], "circle": direct["circle"], "division": ""}
                    elif direct["zone"]:
                        context = {"zone": direct["zone"], "circle": "", "division": ""}
                    else:
                        direct.update(context)
                        if direct["zone"] and direct["circle"]:
                            direct["mappingStatus"] = "resolved"
                        elif direct["zone"] or direct["circle"] or direct["division"]:
                            direct["mappingStatus"] = "partial"
                hierarchy_override = direct

            record = record_from_row(
                row, sheet_name, row_number, resolver, enrichment, args.debug, hierarchy_override
            )
            if record:
                records.append(record)

    ids = [r["id"] for r in records]
    if len(ids) != len(set(ids)):
        raise SystemExit("Duplicate record IDs detected; conversion stopped")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    write_report(args.report, records)

    field = [r for r in records if r["category"] != "Corporate Office"]
    resolved = sum(bool(r["zone"] and r["circle"]) for r in field)
    partial = sum(bool(r["zone"] or r["circle"] or r["division"]) and not (r["zone"] and r["circle"]) for r in field)
    unresolved = len(field) - resolved - partial
    print(f"Wrote {len(records)} contacts to {args.output}")
    print(f"Field hierarchy: {resolved} resolved, {partial} partial, {unresolved} unresolved")
    print(f"Review partial/unresolved rows in {args.report}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
